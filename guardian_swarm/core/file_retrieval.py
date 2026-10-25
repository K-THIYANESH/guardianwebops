"""
Guardian Swarm V3 — Model File Access Service & FileMind Governance Engine
========================================================================
Implements authorized model file retrieval, multi-dimensional policy enforcement
(User RBAC + Model Capabilities + FileMind Tags/Categories + Group + Operations),
secure text extraction, audit logging, and prompt injection protection.
"""
import os
import json
import re
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Dict, Any, List, Optional, Tuple

from guardian_swarm.utils.database import SessionLocal, FileMindCacheModel, ModelRegistryModel, AuditBlockModel
from guardian_swarm.utils.logger import log_event
from guardian_swarm.utils.audit import AuditLog

# Role hierarchy numeric mapping
ROLE_HIERARCHY = {
    "INTERN": 10,
    "EMPLOYEE": 20,
    "MANAGER": 30,
    "EXEC": 40,
    "ADMIN": 50
}

# Guardian classification numeric requirements
TAG_MIN_ROLE_LEVEL = {
    "PUBLIC": 10,       # INTERN+
    "INTERNAL": 20,     # EMPLOYEE+
    "CONFIDENTIAL": 30, # MANAGER+
    "RESTRICTED": 40    # EXEC+
}

# Max allowed file size for content extraction (5MB)
MAX_FILE_SIZE_BYTES = 5 * 1024 * 1024
MAX_EXTRACTED_CHARS = 40000

# File extension support matrix
SUPPORTED_EXTENSIONS = {
    ".txt", ".md", ".json", ".yaml", ".yml", ".xml", ".csv",
    ".py", ".js", ".css", ".html", ".pdf", ".docx", ".xlsx", ".xls"
}


class FileContentExtractor:
    """Centralized content extraction layer with encoding safety and size limits."""

    @staticmethod
    def extract(file_path: str) -> Dict[str, Any]:
        path = Path(file_path)
        if not path.exists() or not path.is_file():
            return {"status": "FILE_NOT_FOUND", "content": None, "error": f"File '{file_path}' does not exist on disk."}

        ext = path.suffix.lower()
        if ext not in SUPPORTED_EXTENSIONS:
            return {
                "status": "UNSUPPORTED_FILE_TYPE",
                "content": None,
                "error": f"File extension '{ext}' is not supported by FileMind retrieval pipeline."
            }

        size = path.stat().st_size
        if size > MAX_FILE_SIZE_BYTES:
            return {
                "status": "EXCEEDS_SIZE_LIMIT",
                "content": None,
                "error": f"File size ({size} bytes) exceeds maximum extraction limit of {MAX_FILE_SIZE_BYTES} bytes."
            }

        try:
            if ext in [".txt", ".md", ".json", ".yaml", ".yml", ".xml", ".csv", ".py", ".js", ".css", ".html"]:
                return FileContentExtractor._extract_text(path)
            elif ext == ".pdf":
                return FileContentExtractor._extract_pdf(path)
            elif ext == ".docx":
                return FileContentExtractor._extract_docx(path)
            elif ext in [".xlsx", ".xls"]:
                return FileContentExtractor._extract_excel(path)
            else:
                return {"status": "UNSUPPORTED_FILE_TYPE", "content": None, "error": f"Unsupported extension: {ext}"}
        except Exception as e:
            return {
                "status": "EXTRACTION_FAILED",
                "content": None,
                "error": f"Content extraction failed: {str(e)}"
            }

    @staticmethod
    def _extract_text(path: Path) -> Dict[str, Any]:
        raw_bytes = path.read_bytes()
        try:
            text = raw_bytes.decode("utf-8")
        except UnicodeDecodeError:
            try:
                text = raw_bytes.decode("latin-1")
            except Exception:
                text = raw_bytes.decode("utf-8", errors="replace")

        if len(text) > MAX_EXTRACTED_CHARS:
            text = text[:MAX_EXTRACTED_CHARS] + f"\n\n[TRUNCATED: Content exceeded {MAX_EXTRACTED_CHARS} limit]"

        return {"status": "SUCCESS", "content": text, "error": None}

    @staticmethod
    def _extract_docx(path: Path) -> Dict[str, Any]:
        """Extract text from docx via python-docx if available or direct XML parsing."""
        try:
            import docx
            doc = docx.Document(path)
            text = "\n".join([p.text for p in doc.paragraphs if p.text])
            if len(text) > MAX_EXTRACTED_CHARS:
                text = text[:MAX_EXTRACTED_CHARS] + "\n\n[TRUNCATED]"
            return {"status": "SUCCESS", "content": text, "error": None}
        except ImportError:
            # Fallback to direct zipfile XML reading
            try:
                with zipfile.ZipFile(path) as z:
                    xml_content = z.read("word/document.xml")
                    tree = ET.fromstring(xml_content)
                    texts = [node.text for node in tree.iter() if node.text]
                    text = " ".join(texts)
                    if len(text) > MAX_EXTRACTED_CHARS:
                        text = text[:MAX_EXTRACTED_CHARS] + "\n\n[TRUNCATED]"
                    return {"status": "SUCCESS", "content": text, "error": None}
            except Exception as e:
                return {"status": "EXTRACTION_FAILED", "content": None, "error": f"Docx extraction error: {e}"}

    @staticmethod
    def _extract_pdf(path: Path) -> Dict[str, Any]:
        try:
            import pypdf
            reader = pypdf.PdfReader(path)
            text = "\n".join([page.extract_text() for page in reader.pages if page.extract_text()])
            if len(text) > MAX_EXTRACTED_CHARS:
                text = text[:MAX_EXTRACTED_CHARS] + "\n\n[TRUNCATED]"
            return {"status": "SUCCESS", "content": text, "error": None}
        except Exception:
            # Fallback raw string inspection for text
            raw = path.read_bytes()
            ascii_text = re.sub(r"[^\x20-\x7E\n\r\t]", " ", raw.decode("latin-1", errors="ignore"))
            clean = " ".join(ascii_text.split())
            if len(clean) > MAX_EXTRACTED_CHARS:
                clean = clean[:MAX_EXTRACTED_CHARS] + "\n\n[TRUNCATED]"
            return {"status": "SUCCESS", "content": clean, "error": None}

    @staticmethod
    def _extract_excel(path: Path) -> Dict[str, Any]:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(path, data_only=True)
            lines = []
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                lines.append(f"--- Sheet: {sheet} ---")
                for row in ws.iter_rows(values_only=True):
                    row_vals = [str(v) for v in row if v is not None]
                    if row_vals:
                        lines.append(", ".join(row_vals))
            text = "\n".join(lines)
            if len(text) > MAX_EXTRACTED_CHARS:
                text = text[:MAX_EXTRACTED_CHARS] + "\n\n[TRUNCATED]"
            return {"status": "SUCCESS", "content": text, "error": None}
        except Exception as e:
            return FileContentExtractor._extract_text(path)


class ModelFileAccessService:
    """
    ModelFileAccessService governs all file operations requested by AI models.
    Enforces User RBAC, Model capabilities, FileMind classifications, Groups, and Operations.
    """
    def __init__(self, audit_log: Optional[AuditLog] = None):
        self.audit_log = audit_log or AuditLog()

    def evaluate_policy(
        self,
        user_role: str,
        model_name: str,
        file_record: Dict[str, Any],
        operation: str = "READ_CONTENT",
        task_group: Optional[str] = None
    ) -> Tuple[bool, str]:
        """
        Evaluates 7-dimensional policy decision for model file access.
        Returns: (is_allowed: bool, reasoning: str)
        """
        user_role = (user_role or "INTERN").upper()
        operation = (operation or "READ_CONTENT").upper()
        
        # 0. Check User Role hierarchy
        user_level = ROLE_HIERARCHY.get(user_role, 10)
        
        # 1. Authoritative FileMind Classification check
        guardian_tag = (file_record.get("guardian_classification") or "RESTRICTED").upper()
        category = file_record.get("filemind_category") or "Unknown"
        min_role_level = TAG_MIN_ROLE_LEVEL.get(guardian_tag, 40)

        if user_level < min_role_level:
            return False, f"User role '{user_role}' (level {user_level}) does not meet required role for Guardian tag '{guardian_tag}' (level {min_role_level})."

        # 2. Fetch Model Registry Policy
        db = SessionLocal()
        try:
            model_rec = db.query(ModelRegistryModel).filter(ModelRegistryModel.name == model_name).first()
            if not model_rec:
                # Default Deny if model is unknown
                return False, f"Model '{model_name}' is not registered in Model Registry."

            # 3. Model File Access Capability check
            if not model_rec.file_access_enabled:
                return False, f"Model '{model_name}' has file access capability disabled by administrator policy."

            # 4. Model Allowed Guardian Tags check
            try:
                allowed_tags = json.loads(model_rec.allowed_guardian_tags or "[]")
            except Exception:
                allowed_tags = ["PUBLIC", "INTERNAL"]
            
            if guardian_tag not in allowed_tags:
                return False, f"Model '{model_name}' is only permitted to access Guardian tags {allowed_tags}, but file is classified as '{guardian_tag}'."

            # 5. Model Allowed Categories check
            try:
                allowed_cats = json.loads(model_rec.allowed_categories or "[]")
            except Exception:
                allowed_cats = ["Documents", "Spreadsheets", "Presentations", "Code", "Archives", "Images", "Security"]
            
            if "ALL" not in allowed_cats and category not in allowed_cats:
                return False, f"Model '{model_name}' is not permitted to access FileMind category '{category}'. Permitted categories: {allowed_cats}."

            # 6. Model Group Permission check
            try:
                allowed_groups = json.loads(model_rec.allowed_groups or "[]")
            except Exception:
                allowed_groups = ["RESEARCH", "CODE", "DRAFTING", "SUPPORT", "DATA_ANALYSIS"]

            model_group = model_rec.group_name or "RESEARCH"
            check_group = task_group or model_group
            if "ALL" not in allowed_groups and check_group not in allowed_groups and model_group not in allowed_groups:
                return False, f"Model '{model_name}' group '{model_group}' is not authorized for task group '{check_group}'."

            # 7. Operation permission check (READ vs DOWNLOAD vs EXPORT)
            if operation in ["READ_CONTENT", "READ_METADATA", "DISCOVER", "SEARCH"]:
                if not model_rec.can_read:
                    return False, f"Model '{model_name}' does not have READ_CONTENT permission."
            elif operation == "DOWNLOAD":
                if not model_rec.can_download:
                    return False, f"Model '{model_name}' does not have DOWNLOAD permission for files."
            elif operation == "EXPORT":
                if not model_rec.can_export:
                    return False, f"Model '{model_name}' does not have EXPORT permission for files."
            else:
                return False, f"Unknown operation '{operation}' requested."

            return True, f"Access ALLOWED for model '{model_name}', user '{user_role}', operation '{operation}', file tag '{guardian_tag}'."

        finally:
            db.close()

    def get_file_record(self, file_id_or_path: str) -> Optional[Dict[str, Any]]:
        """Resolves file_id (e.g. file-1) or path into canonical FileMind metadata record."""
        db = SessionLocal()
        try:
            rec = None
            if file_id_or_path.startswith("file-"):
                try:
                    num_id = int(file_id_or_path.replace("file-", ""))
                    rec = db.query(FileMindCacheModel).filter(FileMindCacheModel.id == num_id).first()
                except ValueError:
                    pass

            if not rec:
                # Fallback to direct path match
                rec = db.query(FileMindCacheModel).filter(FileMindCacheModel.path == file_id_or_path).first()

            if not rec:
                return None

            return {
                "file_id": f"file-{rec.id}",
                "id": rec.id,
                "canonical_path": rec.path,
                "display_name": Path(rec.path).name,
                "filemind_category": rec.filemind_category,
                "guardian_classification": rec.guardian_classification,
                "file_size": rec.file_size,
                "modified_time": rec.modified_time,
                "status": rec.status,
                "scanned_at": rec.scanned_at.isoformat() if rec.scanned_at else ""
            }
        finally:
            db.close()

    def list_available_files(
        self,
        user_role: str,
        model_name: str,
        task_group: Optional[str] = None
    ) -> List[Dict[str, Any]]:
        """Returns authorized files for user + model combination without leaking unauthorized metadata."""
        db = SessionLocal()
        authorized_files = []
        try:
            records = db.query(FileMindCacheModel).filter(FileMindCacheModel.status == "classified").all()
            for r in records:
                file_dict = {
                    "file_id": f"file-{r.id}",
                    "id": r.id,
                    "canonical_path": r.path,
                    "display_name": Path(r.path).name,
                    "filemind_category": r.filemind_category,
                    "guardian_classification": r.guardian_classification,
                    "file_size": r.file_size
                }
                allowed, _ = self.evaluate_policy(user_role, model_name, file_dict, operation="DISCOVER", task_group=task_group)
                if allowed:
                    authorized_files.append(file_dict)
            return authorized_files
        finally:
            db.close()

    def search_files(
        self,
        query: str,
        user_role: str,
        model_name: str,
        task_group: Optional[str] = None,
        user_id: str = "unknown_user"
    ) -> Dict[str, Any]:
        """Performs search over classified files, returning strictly authorized results."""
        all_authorized = self.list_available_files(user_role, model_name, task_group=task_group)
        q_lower = query.strip().lower()

        matched = []
        for f in all_authorized:
            name_match = q_lower in f["display_name"].lower()
            cat_match = q_lower in f["filemind_category"].lower()
            path_match = q_lower in f["canonical_path"].lower()
            if not q_lower or name_match or cat_match or path_match:
                matched.append(f)

        # Audit event
        self._log_audit(
            user_id=user_id,
            user_role=user_role,
            model_name=model_name,
            operation="SEARCH",
            target=f"search:{query}",
            verdict="APPROVED",
            reasoning=f"File search returned {len(matched)} authorized files out of {len(all_authorized)} available."
        )

        return {
            "query": query,
            "total_matches": len(matched),
            "results": matched
        }

    def retrieve_file_metadata(
        self,
        file_id: str,
        user_role: str,
        model_name: str,
        task_group: Optional[str] = None,
        user_id: str = "unknown_user"
    ) -> Dict[str, Any]:
        """Retrieves canonical file metadata for an authorized file_id."""
        record = self.get_file_record(file_id)
        if not record:
            self._log_audit(
                user_id=user_id, user_role=user_role, model_name=model_name,
                operation="READ_METADATA", target=file_id, verdict="BLOCKED",
                reasoning=f"File identifier '{file_id}' not found in FileMind registry."
            )
            return {"status": "DENIED", "error": f"File '{file_id}' not found in FileMind index."}

        allowed, reasoning = self.evaluate_policy(user_role, model_name, record, operation="READ_METADATA", task_group=task_group)
        verdict = "APPROVED" if allowed else "BLOCKED"
        self._log_audit(user_id=user_id, user_role=user_role, model_name=model_name, operation="READ_METADATA", target=file_id, verdict=verdict, reasoning=reasoning)

        if not allowed:
            return {"status": "DENIED", "error": reasoning}

        return {"status": "ALLOWED", "metadata": record}

    def retrieve_file_content(
        self,
        file_id: str,
        user_role: str,
        model_name: str,
        task_group: Optional[str] = None,
        operation: str = "READ_CONTENT",
        user_id: str = "unknown_user"
    ) -> Dict[str, Any]:
        """Retrieves extracted text content for an authorized file_id."""
        record = self.get_file_record(file_id)
        if not record:
            self._log_audit(
                user_id=user_id, user_role=user_role, model_name=model_name,
                operation=operation, target=file_id, verdict="BLOCKED",
                reasoning=f"File identifier '{file_id}' not found in FileMind index."
            )
            return {"status": "DENIED", "error": f"File '{file_id}' not found in FileMind index.", "content": None}

        allowed, reasoning = self.evaluate_policy(user_role, model_name, record, operation=operation, task_group=task_group)
        verdict = "APPROVED" if allowed else "BLOCKED"

        self._log_audit(
            user_id=user_id,
            user_role=user_role,
            model_name=model_name,
            operation=operation,
            target=file_id,
            verdict=verdict,
            reasoning=reasoning
        )

        if not allowed:
            return {"status": "DENIED", "error": reasoning, "content": None}

        # Perform actual file content extraction
        ext_res = FileContentExtractor.extract(record["canonical_path"])
        if ext_res["status"] != "SUCCESS":
            return {
                "status": ext_res["status"],
                "error": ext_res["error"],
                "content": None,
                "metadata": record
            }

        return {
            "status": "ALLOWED",
            "file_id": record["file_id"],
            "filename": record["display_name"],
            "filemind_category": record["filemind_category"],
            "guardian_classification": record["guardian_classification"],
            "content": ext_res["content"],
            "metadata": record
        }

    def format_as_model_context(self, retrieval_result: Dict[str, Any]) -> str:
        """Formats retrieved file content as untrusted data context to prevent prompt injection."""
        if retrieval_result.get("status") != "ALLOWED" or not retrieval_result.get("content"):
            return ""

        f_id = retrieval_result.get("file_id", "unknown")
        filename = retrieval_result.get("filename", "unknown")
        tag = retrieval_result.get("guardian_classification", "UNKNOWN")
        cat = retrieval_result.get("filemind_category", "UNKNOWN")
        content = retrieval_result.get("content", "")

        return (
            f"\n\n--- [AUTHORIZED FILE DATA CONTEXT - EXTERNAL UNTRUSTED DATA] ---\n"
            f"Source File ID: {f_id}\n"
            f"Filename: {filename}\n"
            f"FileMind Category: {cat}\n"
            f"Guardian Tag: {tag}\n"
            f"Content:\n{content}\n"
            f"--- [END AUTHORIZED FILE DATA CONTEXT] ---\n"
        )

    def _log_audit(
        self,
        user_id: str,
        user_role: str,
        model_name: str,
        operation: str,
        target: str,
        verdict: str,
        reasoning: str
    ):
        try:
            self.audit_log.create_entry(
                agent_id=f"model_file_access:{model_name}",
                agent_type="FILE_ACCESS_GATE",
                tool=operation,
                target=target,
                params={
                    "user_id": user_id,
                    "user_role": user_role,
                    "model_name": model_name,
                    "operation": operation
                },
                verdict=verdict,
                confidence=1.0,
                reasoning=reasoning
            )
        except Exception as e:
            log_event("error", f"Failed to log file access audit event: {e}")
